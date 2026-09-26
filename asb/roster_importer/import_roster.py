#!/usr/bin/env python3
"""ASB roster importer (read-only) - Roster 2026 workbook -> normalised events + review report.

Usage:
  python3 import_roster.py <roster.xlsx | drive-download.json> [options]

Options:
  --from YYYY-MM-DD     first date to import (default 2026-06-01)
  --to   YYYY-MM-DD     last date to import  (default 2026-09-30)
  --aliases FILE        optional CSV of confirmed name fixes: raw,canonical,role
                        (role = coach | athlete). Nothing is merged unless it is in
                        this file or in COACHES below.
  --out DIR             output folder (default: roster_import_out)

What it reads (never writes) - the "Roster 2026" Google Sheet exported as .xlsx:
  * month tabs made of weekly blocks: a header row of dates in columns B,D,F,H,J,L,N,
    30-minute rows (time label in column A; per day a session column and a coach column),
    then one row per coach with hours or WO / L / SL, and a weekly total after Sunday;
  * PT-group / squad side-tables to the right of the grid (coach row, group row, players).

What it writes, all inside --out:
  report.html          the review report (open on a phone or laptop)
  events.csv           one row per event (date, start, end, type, title, coaches, athletes ...)
  issues.csv           everything the importer could not resolve, with the sheet cell
  names.csv            every raw name spelling seen and what it resolved to
  memberships.csv      athlete/coach -> group seed from the side-tables
  staff_status.csv     per coach per day: hours or WO / L / SL
  roster_import.json   the full normalised model (person, group, membership, event,
                       event_group, event_person, staff_status, import_issue)

Nothing is guessed silently: every ambiguity becomes a row in issues.csv.
Only Python 3.9+ and openpyxl are needed.
"""
import argparse
import base64
import collections
import csv
import datetime as dt
import hashlib
import html
import io
import json
import os
import re
import sys

try:
    import openpyxl
except ImportError:  # pragma: no cover
    sys.exit("openpyxl is missing. Run:  python3 -m pip install openpyxl")

SHEET_URL = "https://docs.google.com/spreadsheets/d/1kKNIbgWmPDmW0E_4gLRW1fEqhMX3VSZIiDKxKeIBo3c/edit"
DAY_COLS = [2, 4, 6, 8, 10, 12, 14]          # B D F H J L N (session col; coach col = +1)
TOTAL_COL = 16                               # P - weekly total on the coach rows
SIDE_FIRST_COL = 17                          # Q - side-tables start at or after this column
MONTHS = ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"]
# Only month tabs ("Jan ", "March ", "Sep") and month copies ("June 26") hold weekly grids.
# Every other tab (S&C, Physio, individual plans, scratch tabs) is listed in the report as
# not imported - a later phase.

# ---------------------------------------------------------------- vocabulary
# Coaches and staff. Key = canonical name; value = confirmed spellings (lower case,
# honorifics like "sir", "di", "ma'am" are stripped before matching).
COACHES = {
    "Manoj": ["manoj"],
    "Puneeth": ["puneeth", "puneet"],
    "Rudra": ["rudra", "rura"],
    "Aditya": ["aditya"],
    "Suraj": ["suraj", "surajj"],
    "Mario": ["mario"],
    "Divya": ["divya"],
    "Sahil": ["sahil"],
    "Sanket": ["sanket"],
    "Anu": ["anu"],
    "Guru": ["guru"],
    "Rishi": ["rishi"],
    "Nityn": ["nityn", "nityan", "nityan b", "nityn b"],
}
ALL_COACHES = "All coaches"
ALL_SPELLINGS = {"all", "all coaches", "everyone"}
HONORIFICS = r"\b(?:sir|didi|di|ma'?am|maam|mam|bhaiya)\b"
STATUS = {"WO": "WO", "W/O": "WO", "WEEKOFF": "WO", "L": "L", "LEAVE": "L", "SL": "SL",
          "CL": "L", "H": "holiday", "HOLIDAY": "holiday", "COMP": "comp_off", "CO": "comp_off"}

# Groups: (canonical name, kind, regex). Order matters - longer phrases first.
GROUP_PATTERNS = [
    ("SR & U18 Elite", "squad", r"sr\s*(?:and|&)\s*u18\s*elite"),
    ("ASB Jr Elite", "squad", r"asb\s*jr\.?\s*elite"),
    ("SC Elite", "camp", r"sc\s*elite"),
    ("SC Beg", "camp", r"sc\s*beg(?:inner)?"),
    ("SC Mini", "camp", r"sc\s*mini"),
    ("SC 3x3", "camp", r"sc\s*3\s*x\s*3"),
    ("Jr Elite", "squad", r"(?:jr|junior)\.?\s*elite"),
    ("ASB Seniors", "squad", r"asb\s*(?:seniors?|sr)|\bsr\s+asb|\bseniors\b"),
    ("ASB Juniors", "squad", r"asb\s*(?:juniors?|jr)|\bjr\s+asb"),
    ("ASB U18", "team", r"(?:asb\s*)?\bu\s?18\b"),
    ("ASB Rehab", "squad", r"asb\s*rehab"),
    ("Elite", "squad", r"\belite\b"),
    ("Intermediate", "class", r"\bintermediate\b|\bint\b"),
    ("Beginner", "class", r"\bbeginners?\b|\bbeg\b"),
    ("Mini", "class", r"\bmini\b"),
    ("JC Pro", "partner", r"\bjc\s*pro\b"),
    ("CSE", "partner", r"\bcse\b"),
    ("Oyme", "partner", r"\boyme\b"),
    ("Sadhu Vaswani", "partner", r"sadhu\s*vaswani"),
    ("Delhi team", "partner", r"\bdelhi(?:\s*(?:boys|team))?\b"),
    ("Mumbai kids", "partner", r"mumbai\s*kids"),
    ("Akshar", "partner", r"\bakshar\b"),
    ("Residential athletes", "squad", r"residential\s*athletes"),
    ("Small PT", "pt_group", r"small\s*pt"),
    ("Girls PT", "pt_group", r"girls\s*pt"),
    ("PT Rehab", "pt_group", r"pt\s*rehab"),
    ("SC", "camp", r"\bsc\b"),
]
# Activities: (label, event type, regex). The first matching type in TYPE_ORDER wins.
ACTIVITIES = [
    ("Holiday", "holiday", r"\bholiday\b"),
    ("Internal tournament", "tournament", r"internal\s*tour(?:nament)?"),
    ("3x3 tournament", "tournament", r"3\s*x\s*3\s*tournament"),
    ("Tournament", "tournament", r"\btournament\b"),
    ("Ice bath", "ice_bath", r"ice\s*bath"),
    ("Coaches meeting", "meeting", r"coach(?:es)?\s*meeting|\bmeeting\b"),
    ("Film session", "review", r"\bfilm(?:\s*(?:session|study))?\b"),
    ("Match", "match", r"\bmatch\b|\bgame\b|\bvs\.?\b|\b3\s*x\s*3\b|\b5\s*x\s*5\b"),
    ("Physio", "physio", r"\bphysio\b"),
    ("S&C", "sc", r"\bs\s*&\s*c\b|\bsnc\b|\bgym\b"),
    ("Assessment", "other", r"\bassessment\b|shooting\s*test"),
    ("Shooting", "session", r"\bshoot(?:ing)?\b"),
    ("Scrimmage", "session", r"\bscrimmage\b"),
    ("Contest", "session", r"\bcontest\b"),
    ("Closing", "session", r"\bclosing\b"),
    ("Stations", "session", r"\bstations?\b"),
    ("Classroom", "session", r"\bclassroom\b"),
    ("Yoga", "session", r"\byoga\b"),
    ("Volleyball", "session", r"\bvoll?[ey]*ball\b"),
    ("Swimming", "session", r"\bswimming\b"),
    ("Orientation", "other", r"\bori[ea]ntation\b"),
    ("Photoshoot", "other", r"\bphoto\s*shoot\b"),
    ("Cleaning", "other", r"\bcleaning\b"),
    ("Recovery / mobility", "other", r"\brecovery\b|\bmobility\b"),
    ("Event", "other", r"\bevent\b"),
]
TYPE_ORDER = ["holiday", "tournament", "match", "ice_bath", "meeting", "review", "physio", "sc", "other"]
COACHLESS_TYPES = {"ice_bath", "holiday"}
# Words that are not names when left over after groups/activities are removed.
FRAGMENTS = {"jr", "sr", "bb", "vs", "sc", "pt"}
STOPWORDS = {"pt", "session", "sessions", "the", "and", "with", "for", "group", "grp", "cls",
             "class", "rest", "on", "court", "only", "test", "extra", "practice", "training"}

# ---------------------------------------------------------------- small helpers
def norm_space(s):
    return re.sub(r"\s+", " ", str(s)).strip()


def slug(s):
    return re.sub(r"[^a-z0-9]+", "-", str(s).lower()).strip("-") or "x"


def nice_case(name):
    name = norm_space(name)
    if name.islower() or name.isupper():
        return " ".join(w[:1].upper() + w[1:].lower() for w in name.split(" "))
    return name[:1].upper() + name[1:]


def levenshtein(a, b):
    if a == b:
        return 0
    prev = list(range(len(b) + 1))
    for i, ca in enumerate(a, 1):
        cur = [i]
        for j, cb in enumerate(b, 1):
            cur.append(min(prev[j] + 1, cur[j - 1] + 1, prev[j - 1] + (ca != cb)))
        prev = cur
    return prev[-1]


def col_letter(c):
    s = ""
    while c:
        c, r = divmod(c - 1, 26)
        s = chr(65 + r) + s
    return s


def cell_ref(tab, row, col, col2=None, row2=None):
    a = f"{col_letter(col)}{row}"
    if col2 or row2:
        a += f":{col_letter(col2 or col)}{row2 or row}"
    return f"'{tab.strip()}'!{a}"


def fmt_time(mins):
    h, m = divmod(mins, 60)
    return f"{h:02d}:{m:02d}"


def fmt_12(mins):
    h, m = divmod(mins, 60)
    ap = "am" if h < 12 else "pm"
    h12 = h % 12 or 12
    return f"{h12}:{m:02d} {ap}" if m else f"{h12} {ap}"


# ---------------------------------------------------------------- loading
def load_workbook(path):
    with open(path, "rb") as f:
        raw = f.read()
    if raw[:2] != b"PK":                      # Google Drive download wrapped in JSON
        raw = base64.b64decode(json.loads(raw.decode("utf-8"))["content"])
    return openpyxl.load_workbook(io.BytesIO(raw), data_only=True, read_only=False)


class Grid:
    """Cell access with merged ranges filled from their top-left cell."""

    def __init__(self, ws):
        self.ws = ws
        self.title = ws.title
        self.origin = {}
        for rng in ws.merged_cells.ranges:
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    self.origin[(r, c)] = (rng.min_row, rng.min_col)
        self.max_row = ws.max_row
        self.max_col = ws.max_column

    def get(self, r, c):
        rr, cc = self.origin.get((r, c), (r, c))
        return self.ws.cell(rr, cc).value

    def text(self, r, c):
        v = self.get(r, c)
        if v is None:
            return ""
        if isinstance(v, (dt.datetime, dt.date)):
            return v.strftime("%Y-%m-%d")
        return str(v)

    def same_merge(self, a, b):
        oa, ob = self.origin.get(a), self.origin.get(b)
        return oa is not None and oa == ob


def is_month_tab(title):
    t = title.strip().lower()
    return bool(re.fullmatch(r"[a-z]+", t)) and t[:3] in MONTHS and len(t) <= 9


def is_month_copy(title):
    """'June 26' - a copy of a month tab. Read, but the real month tab wins on clashes."""
    m = re.fullmatch(r"([a-z]+)\s+\d{1,4}", title.strip().lower())
    return bool(m) and m.group(1)[:3] in MONTHS


def as_date(v, year=2026):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, str):
        m = re.fullmatch(r"\s*(\d{1,2})\s*(?:st|nd|rd|th)?[\s\-/]*([A-Za-z]{3,9})\.?\s*(\d{4})?\s*", v)
        if m and m.group(2)[:3].lower() in MONTHS:
            try:
                return dt.date(int(m.group(3) or year), MONTHS.index(m.group(2)[:3].lower()) + 1, int(m.group(1)))
            except ValueError:
                return None
    return None


# ---------------------------------------------------------------- time labels
TIME_TOKEN = re.compile(r"(\d{1,2})(?:\s*[:.]\s*(\d{1,2})?)?\s*([ap])?\.?\s*m?\.?", re.I)


def parse_time_label(label):
    """'8:00AM- 8:30AM', '12:30PM - 1.00PM', '9:30AM - 10:AM' -> (start_min, end_min) or None."""
    if not isinstance(label, str) or not re.search(r"\d", label):
        return None
    toks = []
    for m in TIME_TOKEN.finditer(label):
        h = int(m.group(1))
        mi = int(m.group(2)) if m.group(2) else 0
        if h > 23 or mi > 59:
            return None
        toks.append([h, mi, (m.group(3) or "").lower()])
    if not toks:
        return None
    if len(toks) >= 2 and not toks[0][2]:
        toks[0][2] = toks[1][2]

    def to_min(h, mi, ap):
        if ap == "p" and h != 12:
            h += 12
        elif ap == "a" and h == 12:
            h = 0
        elif not ap and h < 7:                # un-marked afternoon labels ("1:00 - 1:30")
            h += 12
        return h * 60 + mi

    start = to_min(*toks[0])
    end = to_min(*toks[1]) if len(toks) >= 2 else start + 30
    if end <= start:
        end = start + 30
    if not (6 * 60 <= start <= 22 * 60):
        return None
    return start, end


# ---------------------------------------------------------------- issues
class Issues:
    def __init__(self):
        self.rows = []
        self._seen = set()

    def add(self, severity, kind, where, raw, problem, suggestion="", date=None):
        key = (kind, where, raw, problem)
        if key in self._seen:
            return
        self._seen.add(key)
        self.rows.append({"severity": severity, "kind": kind, "date": str(date or ""), "cell": where,
                          "raw": " ⏎ ".join(norm_space(x) for x in str(raw).split("\n") if norm_space(x))[:300], "problem": problem, "suggestion": suggestion})


# ---------------------------------------------------------------- names
class Names:
    """Resolves raw spellings to people. Only confirmed spellings are merged."""

    def __init__(self, aliases_path=None):
        self.coach_alias = {}
        for canon, spellings in COACHES.items():
            for s in spellings + [canon.lower()]:
                self.coach_alias[s] = canon
        self.athlete_alias = {}
        if aliases_path:
            with open(aliases_path, newline="", encoding="utf-8") as f:
                for row in csv.DictReader(f):
                    raw = norm_space(row.get("raw", "")).lower()
                    canon = norm_space(row.get("canonical", ""))
                    role = (row.get("role") or "athlete").strip().lower()
                    if not raw or not canon:
                        continue
                    if role == "coach":
                        self.coach_alias[raw] = canon
                    else:
                        self.athlete_alias[raw] = canon
        self.seen = collections.Counter()   # (raw, resolved, role) -> count
        self.athletes = {}                  # canonical -> {"first_seen", "sources"}

    # coaches ---------------------------------------------------------
    def coach(self, token):
        """Return canonical coach name or None for one person token."""
        t = re.sub(HONORIFICS, " ", token.lower())
        t = norm_space(re.sub(r"[^a-z' ]", " ", t))
        if not t:
            return ""
        if t in ALL_SPELLINGS:
            return ALL_COACHES
        return self.coach_alias.get(t)

    def coach_suggestion(self, token):
        t = norm_space(re.sub(r"[^a-z ]", " ", re.sub(HONORIFICS, " ", token.lower())))
        best = sorted((levenshtein(t, a), c) for a, c in self.coach_alias.items() if abs(len(a) - len(t)) <= 2)
        best = [c for d, c in best if d <= 2 and d == best[0][0]] if best else []
        return sorted(set(best))

    def split_joined(self, token):
        """'manojsanket' -> ['Manoj', 'Sanket'] when it is exactly two or more known names."""
        t = re.sub(r"[^a-z]", "", token.lower())
        keys = sorted({k for k in self.coach_alias if " " not in k and len(k) >= 3}, key=len, reverse=True)

        def rec(s):
            if not s:
                return []
            for k in keys:
                if s.startswith(k):
                    rest = rec(s[len(k):])
                    if rest is not None:
                        return [self.coach_alias[k]] + rest
            return None

        out = rec(t)
        return out if out and len(out) > 1 else None

    def parse_coach_text(self, text):
        """One coach 'line' -> (resolved names, unresolved raw tokens, had_junk)."""
        s = str(text)
        junk = bool(re.search(r"[*$#+]{2,}|#name\?|#error!|\.{2,}", s, re.I))
        # "Puneeth(Nityan B.)" -> "Puneeth, Nityan B"
        s = re.sub(r"\(([^)]*)\)", r", \1", s)
        s = re.sub(r"#name\?|#error!", " ", s, flags=re.I)
        s = re.sub(r"[*$#.]+", " ", s)
        s = re.sub(HONORIFICS, " ", s, flags=re.I)
        resolved, unresolved = [], []
        for chunk in re.split(r"\s*(?:,|/|&|\+|\band\b)\s*", s, flags=re.I):
            chunk = norm_space(chunk)
            if not chunk or not re.search(r"[A-Za-z]", chunk):
                continue
            c = self.coach(chunk)
            if c:
                resolved.append(c)
                self.seen[(chunk, c, "coach")] += 1
                continue
            words = chunk.split(" ")
            per_word = [self.coach(w) for w in words]
            if len(words) > 1 and all(per_word):
                for w, c in zip(words, per_word):
                    resolved.append(c)
                    self.seen[(w, c, "coach")] += 1
                continue
            if len(words) > 1 and any(per_word):
                for w, c in zip(words, per_word):
                    if c:
                        resolved.append(c)
                        self.seen[(w, c, "coach")] += 1
                    else:
                        unresolved.append(w)
                        self.seen[(w, "", "coach")] += 1
                continue
            joined = self.split_joined(chunk)
            if joined:
                resolved += joined
                self.seen[(chunk, " + ".join(joined), "coach")] += 1
                continue
            unresolved.append(chunk)
            self.seen[(chunk, "", "coach")] += 1
        out = []
        for r in resolved:
            if r not in out:
                out.append(r)
        return out, unresolved, junk

    # athletes --------------------------------------------------------
    def athlete(self, raw, source):
        name = nice_case(raw)
        canon = self.athlete_alias.get(name.lower(), name)
        self.seen[(name, canon, "athlete")] += 1
        rec = self.athletes.setdefault(canon, {"sources": collections.Counter(), "spellings": collections.Counter()})
        rec["sources"][source] += 1
        rec["spellings"][name] += 1
        return canon

    def similar_athletes(self):
        """Clusters of athlete names that might be one person. Reported, never merged."""
        names = sorted(self.athletes)
        clusters, used = [], set()
        for i, a in enumerate(names):
            if a in used:
                continue
            group = [a]
            ka = re.sub(r"[^a-z]", "", a.lower())
            fa = a.split(" ")[0].lower()
            for b in names[i + 1:]:
                if b in used:
                    continue
                kb = re.sub(r"[^a-z]", "", b.lower())
                fb = b.split(" ")[0].lower()
                dist = levenshtein(ka, kb)
                close = (min(len(ka), len(kb)) >= 3 and ka[:2] == kb[:2]
                         and (dist <= 1 or (dist == 2 and min(len(ka), len(kb)) >= 6)))
                same_first = fa == fb and len(fa) >= 3
                if close or same_first:
                    group.append(b)
            if len(group) > 1:
                used.update(group)
                clusters.append(group)
        return clusters


# ---------------------------------------------------------------- session text
class ParsedItem:
    __slots__ = ("raw", "type", "title", "groups", "activities", "athletes", "notes", "leftover", "pt")

    def __init__(self, raw):
        self.raw = raw
        self.type = None
        self.title = ""
        self.groups = []
        self.activities = []
        self.athletes = []
        self.notes = []
        self.leftover = []
        self.pt = False


PT_GROUP_RE = re.compile(r"\b(sr|jr)\s*(?:pt\s*)?(\d(?:\s*(?:&|,|and)\s*\d)*)(?!\s*x)", re.I)


def group_kind(name):
    for canon, kind, _ in GROUP_PATTERNS:
        if canon == name:
            return kind
    return "pt_group" if name.startswith("PT ") else "other"


def parse_session_item(text, name_words):
    """Classify one session item, e.g. 'PT Jr 1 & PT Jr 2', 'Asha PT', 'Ice Bath - SC'."""
    item = ParsedItem(norm_space(text))
    # runs of 3+ spaces separate things ("Riya           Dohun"); keep them as separators
    s = " " + " | ".join(norm_space(p) for p in split_wide(str(text).replace("\n", "   "))).lower() + " "
    # parenthetical: keep as text when it is vocabulary, else keep as a note
    def paren(m):
        inner = m.group(1).strip()
        vocab = any(re.search(rx, inner, re.I) for _, _, rx in GROUP_PATTERNS + [(0, 0, r[2]) for r in ACTIVITIES])
        if vocab:
            return f" {inner} "
        if inner:
            item.notes.append(inner)
        return " | "
    s = re.sub(r"\(([^)]*)\)?", paren, s)
    s = s.replace('"', " ")
    found = []                               # (position in text, group name) - keeps written order

    def mask(m):                             # blank out a match without shifting positions
        return "|" + " " * (len(m.group(0)) - 1)

    # PT groups: "PT Sr 3 & PT Sr 4", "Pt Jr 1&4", "JR 3&4 PT", "pt rehab & sr1 & sr2"
    if re.search(r"pt\b|\bpt|\bp\.t\b", s):
        item.pt = True
        for m in PT_GROUP_RE.finditer(s):
            for k, d in enumerate(re.findall(r"\d", m.group(2))):
                found.append((m.start() + k / 10, f"PT {m.group(1).title()} {d}"))
        s = PT_GROUP_RE.sub(mask, s)
    for canon, kind, rx in GROUP_PATTERNS:
        for m in re.finditer(rx, s, re.I):
            found.append((m.start(), canon))
        s = re.sub(rx, mask, s, flags=re.I)
    for pos, g in sorted(found):
        if g not in item.groups:
            item.groups.append(g)
    acts = []
    for label, etype, rx in ACTIVITIES:
        m = re.search(rx, s, re.I)
        if m:
            acts.append((m.start(), label, etype))
            s = re.sub(rx, mask, s, flags=re.I)
    item.activities = [(label, etype) for _, label, etype in sorted(acts)]
    s = re.sub(r"(?<![a-z])p\.?t(?![a-z])", mask, s)   # stray "PT" words
    s = re.sub(r"(?<=[a-z])pt\b", mask, s)              # "4pt", "bbpt"
    # leftovers -> names or junk
    for chunk in re.split(r"\s*(?:\||,|&|\+|/|\band\b|\s-\s|:)\s*", s):
        chunk = norm_space(chunk.strip(" -.*'\""))
        if not chunk:
            continue
        words = [w for w in chunk.split(" ") if w.lower() not in STOPWORDS]
        frags = [w for w in words if len(w) == 1 or w.lower() in FRAGMENTS]
        if frags and len(words) > 1 and len(frags) < len(words) and all(len(w) > 1 for w in frags):
            item.leftover += frags               # "jr asha" -> leftover "jr", name "asha"
            words = [w for w in words if w not in frags]
        if not words:
            continue
        chunk = " ".join(words)
        if not re.fullmatch(r"[A-Za-z][A-Za-z.' ]*", chunk):
            item.leftover.append(chunk)
            continue
        if len(words) == 1 and (len(words[0]) == 1 or words[0].lower() in FRAGMENTS):
            item.leftover.append(chunk)          # "k", "bb", a lone "jr" - fragments, not names
            continue
        # "asha ravi meera" -> three people when each word is a name seen on its own
        if len(words) > 1 and all(w.lower() in name_words for w in words):
            item.athletes += words
        elif len(words) > 3:
            item.leftover.append(chunk)
        else:
            item.athletes.append(chunk)
    # type
    types = [t for _, t in item.activities]
    for t in TYPE_ORDER:
        if t in types:
            item.type = t
            break
    if not item.type:
        if item.pt or any(group_kind(g) == "pt_group" for g in item.groups):
            item.type = "pt"
        elif item.groups or types:
            item.type = "session"
        elif item.athletes:
            item.type = "pt"                  # a bare player name = individual PT
        else:
            item.type = "other"
    # title
    parts = []
    act_labels = [a for a, _ in item.activities]
    if item.type == "pt" and not item.groups and item.athletes:
        parts.append("PT")
    parts += act_labels[:1]
    if item.groups:
        parts.append(" & ".join(item.groups))
    parts += act_labels[1:]
    title = " - ".join(parts) if parts else ""
    if item.athletes:
        who = ", ".join(nice_case(a) for a in item.athletes)
        title = f"{title}: {who}" if title else who
    if item.notes:
        title += f" ({'; '.join(item.notes)})"
    item.title = title or item.raw
    return item


def split_wide(line):
    """Split one line on runs of 3+ spaces or ' / ' (the sheet's other separators)."""
    return [p.strip() for p in re.split(r"\s{3,}|\s/\s", line) if p.strip()]


def nonempty_lines(v):
    if v is None:
        return []
    return [x.strip() for x in str(v).split("\n") if x.strip()]


def raw_lines(v):
    """Lines of a cell, blanks kept (they align sessions with coaches), trailing blanks dropped."""
    if v is None:
        return []
    out = [x.strip() for x in str(v).split("\n")]
    while out and not out[-1]:
        out.pop()
    return out


def is_all_coaches(coach_text):
    t = norm_space(re.sub(r"[^A-Za-z ]", " ", str(coach_text or ""))).lower()
    return t in ALL_SPELLINGS


def pair_cell(session_text, coach_text, classify, is_coach=None):
    """Pair session items with coach lines. Returns (pairs, how).

    pairs = [(session_item_text, coach_line_text or None)].
    how   = "ok"
          | "merged"     an activity line such as "Shooting" was joined to its group line
          | "name_order" a line like "Mario Aditya Suraj" was split into one coach per session
          | "ambiguous"  every coach attached to every session; the caller logs an issue.
    """
    S = nonempty_lines(session_text)
    C = nonempty_lines(coach_text)
    if not S:
        return [], "ok"
    if not C:
        return [(p, None) for line in S for p in split_wide(line)], "ok"
    if is_all_coaches(coach_text):
        return [(p, ALL_COACHES) for line in S for p in split_wide(line)], "ok"
    res = _pair_lines(S, C, raw_lines(session_text), raw_lines(coach_text), classify)
    if res is not None:
        return res, "ok"
    merged = merge_activity_lines(S, classify)
    if len(merged) != len(S):
        res = _pair_lines(merged, C, None, None, classify)
        if res is not None:
            return res, "merged"
    if is_coach:
        people = split_people(C, is_coach)
        if len(people) != len(C):
            for lines in (S, merged):
                items = [p for line in lines for p in split_wide(line)]
                if len(items) == len(people):
                    return list(zip(items, people)), "name_order"
    items = [p for line in S for p in split_wide(line)]
    return [(it, ", ".join(C)) for it in items], "ambiguous"


def _pair_lines(S, C, rawS, rawC, classify):
    if len(S) == len(C):
        out = []
        for s, c in zip(S, C):
            sub = _pair_line(s, c, classify)
            if sub is None:
                return None
            out += sub
        return out
    if rawS and rawC and len(rawS) == len(rawC) and len(rawS) > 1:
        out = []
        for s, c in zip(rawS, rawC):
            if not s:
                if c:
                    return None               # a coach with no session on that line
                continue
            sub = _pair_line(s, c, classify) if c else [(p, None) for p in split_wide(s)]
            if sub is None:
                return None
            out += sub
        return out
    SW = [p for line in S for p in split_wide(line)]
    CW = [p for line in C for p in split_wide(line)]
    if len(SW) == len(CW):
        return list(zip(SW, CW))
    if len(SW) == 1:
        return [(SW[0], ", ".join(C))]
    return assign(SW, C, classify)


def _pair_line(s, c, classify):
    sw, cw = split_wide(s), split_wide(c)
    if len(sw) == len(cw):
        return list(zip(sw, cw))
    if len(sw) == 1:
        return [(s, ", ".join(cw))]
    return assign(sw, [c], classify)


def assign(items, coach_lines, classify):
    """Leave coachless items (ice bath, holiday) without a coach, then retry pairing."""
    need = [i for i, it in enumerate(items) if classify(it).type not in COACHLESS_TYPES]
    if not (len(need) == len(coach_lines) or (len(need) == 1 and coach_lines)):
        return None
    lines = coach_lines if len(need) == len(coach_lines) else [", ".join(coach_lines)]
    out, k = [], 0
    for i, it in enumerate(items):
        if i in need:
            out.append((it, lines[k]))
            k += 1
        else:
            out.append((it, None))
    return out


def merge_activity_lines(lines, classify):
    """['SC & Jr Elite', 'Shooting'] -> ['SC & Jr Elite Shooting']: an activity-only line joins
    the group line above it (or, failing that, the group line below it)."""
    def activity_only(x):
        it = classify(x)
        return bool(it.activities) and not it.groups and not it.athletes and not it.leftover

    def plain_group(x):
        it = classify(x)
        return bool(it.groups) and not it.activities

    out, carry = [], None
    for line in lines:
        if activity_only(line):
            if out and plain_group(out[-1]) and not activity_only(out[-1]):
                out[-1] = f"{out[-1]} {line}"
                continue
            if carry is None:
                carry = line
                continue
        if carry is not None:
            if plain_group(line):
                line = f"{carry} {line}"
            else:
                out.append(carry)
            carry = None
        out.append(line)
    if carry is not None:
        out.append(carry)
    return out


def split_people(coach_lines, is_coach):
    """'Mario Aditya Suraj' (no commas, every word a known coach) -> three entries."""
    out = []
    for line in coach_lines:
        for part in split_wide(line):
            clean = norm_space(re.sub(HONORIFICS, " ", part, flags=re.I))
            words = clean.split(" ")
            if not re.search(r"[,&/+]", clean) and len(words) > 1 and all(is_coach(w) for w in words):
                out += words
            else:
                out.append(part)
    return out


# ---------------------------------------------------------------- importer
class Importer:
    def __init__(self, wb, date_from, date_to, names, issues):
        self.wb = wb
        self.date_from = date_from
        self.date_to = date_to
        self.names = names
        self.issues = issues
        self.grids = {}
        self.slots = []              # raw per-slot items before merging
        self.staff = {}              # (coach, date) -> record
        self.memberships = []        # snapshot rows
        self.skipped_tabs = []
        self.blocks = []             # (grid, header_row, end_row, {col: date}, total_col)
        self._pending = []
        self.day_source = {}         # date -> (grid, header_row, col)
        self.name_words = set()

    # -- structure ------------------------------------------------------
    def grid(self, ws):
        if ws.title not in self.grids:
            self.grids[ws.title] = Grid(ws)
        return self.grids[ws.title]

    def find_blocks(self):
        for ws in self.wb.worksheets:
            if not (is_month_tab(ws.title) or is_month_copy(ws.title)):
                self.skipped_tabs.append((ws.title, "not a month tab (S&C, physio, individual plan or "
                                                    "scratch layout) - later phase"))
                continue
            if ws.max_row < 3:                   # empty tabs (Oct-Dec)
                continue
            g = self.grid(ws)
            headers = []
            for r in range(1, g.max_row + 1):
                typed = {c: as_date(g.get(r, c)) for c in DAY_COLS
                         if isinstance(g.get(r, c), (dt.date, dt.datetime))}
                if len(typed) >= 2:
                    headers.append((r, typed))
            prev_base = None
            for i, (r, typed) in enumerate(headers):
                end = headers[i + 1][0] - 1 if i + 1 < len(headers) else g.max_row
                dates, prev_base = self.fix_header_dates(g, r, end, typed, prev_base)
                if not dates:
                    continue
                if max(dates.values()) < self.date_from or min(dates.values()) > self.date_to:
                    continue
                for issue in self._pending:
                    self.issues.add(*issue[:5], **issue[5])
                last_col = max(dates)
                self.blocks.append((g, r, end, dates, last_col + 2))

    def fix_header_dates(self, g, r, end, typed, prev_base):
        """Dates run one per day across B..N. A copy-pasted header (e.g. last week's dates)
        is corrected from the previous block (+7 days) or, failing that, a majority vote.
        Issues are held in self._pending and only logged for blocks inside the date range."""
        self._pending = []
        bases = collections.Counter(d - dt.timedelta(days=DAY_COLS.index(c)) for c, d in typed.items())
        if prev_base and (prev_base + dt.timedelta(days=7)) in bases:
            base = prev_base + dt.timedelta(days=7)
        else:
            base, votes = bases.most_common(1)[0]
            if votes < 2:
                self._pending.append(("fix", "header_date", cell_ref(g.title, r, 2, 15), str(typed),
                                      "Week header dates are not consecutive; block skipped", {}))
                return {}, prev_base
        first, last = min(typed), max(typed)
        dates = {}
        for i, c in enumerate(DAY_COLS):
            want = base + dt.timedelta(days=i)
            got = typed.get(c)
            if got is None:
                if first < c < last:              # a gap between dated columns
                    dates[c] = want
                    self._pending.append(("check", "header_date", cell_ref(g.title, r, c), "",
                                          f"Header date missing; assumed {want:%a %d %b} from the rest of the row",
                                          {"date": want}))
                continue
            if got != want:
                self._pending.append(("fix", "header_date", cell_ref(g.title, r, c), f"{got:%a %d %b %Y}",
                                      f"Header date looks wrong; used {want:%a %d %b} "
                                      f"(the week runs {base:%d %b} to {base + dt.timedelta(days=6):%d %b})",
                                      {"suggestion": f"Change the header to {want:%d %b %Y}", "date": want}))
            dates[c] = want
        return dates, base

    def choose_day_sources(self):
        cands = collections.defaultdict(list)
        for g, r, end, dates, _ in self.blocks:
            for c, d in dates.items():
                if not (self.date_from <= d <= self.date_to):
                    continue
                filled = sum(1 for rr in range(r + 1, min(end, r + 40) + 1) if g.text(rr, c).strip() or g.text(rr, c + 1).strip())
                cands[d].append((g, r, c, filled))
        for d, lst in sorted(cands.items()):
            month_tab = [x for x in lst if is_month_tab(x[0].title) and x[3] > 0]
            pool = month_tab or [x for x in lst if x[3] > 0] or lst
            pool.sort(key=lambda x: (-(x[0].title.strip().lower()[:3] == MONTHS[d.month - 1]), -x[3]))
            g, r, c, _ = pool[0]
            self.day_source[d] = (g, r, c)
            others = [x for x in lst if x[0] is not g or x[1] != r]
            others = [x for x in others if x[3] > 0]
            if others:
                self.issues.add("check", "duplicate_day", cell_ref(g.title, r, c), "",
                                f"{d:%a %d %b} also appears in " +
                                ", ".join(f"'{x[0].title.strip()}' row {x[1]}" for x in others) +
                                f"; used '{g.title.strip()}' row {r}",
                                suggestion="Delete or rename the stale copy", date=d)

    # -- grid -------------------------------------------------------------
    def scan_names(self):
        """First pass: single-word names seen anywhere, so 'Asha Ravi' can split."""
        for d, (g, r, c) in self.day_source.items():
            for rr in range(r + 1, r + 40):
                v = g.get(rr, c)
                if not v or not parse_time_label(g.get(rr, 1)):
                    continue
                for line in nonempty_lines(v):
                    line = re.sub(r"\([^)]*\)?", " ", line)
                    for part in re.split(r"\s{3,}|,|&|/|\band\b", line, flags=re.I):
                        w = norm_space(part)
                        if re.fullmatch(r"[A-Za-z]{2,}", w or "") and w.lower() not in FRAGMENTS:
                            self.name_words.add(w.lower())
        vocab = {w for w in self.name_words if parse_session_item(w, set()).groups
                 or parse_session_item(w, set()).activities}
        self.name_words -= vocab | STOPWORDS | set(self.names.coach_alias)

    def classify(self, text):
        return parse_session_item(text, self.name_words)

    def read_days(self):
        for d, (g, hdr, col) in sorted(self.day_source.items()):
            end = next((e for gg, r, e, _, _ in self.blocks if gg is g and r == hdr), g.max_row)
            time_rows = []
            staff_rows = []
            for r in range(hdr + 1, end + 1):
                a = g.get(r, 1)
                t = parse_time_label(a) if isinstance(a, str) else None
                if t:
                    time_rows.append((r, t))
                elif isinstance(a, str) and a.strip() and time_rows:
                    staff_rows.append(r)
                elif a is None and time_rows and not staff_rows:
                    if g.get(r, col) or g.get(r, col + 1):
                        self.issues.add("check", "no_time_label", cell_ref(g.title, r, col, col + 1),
                                        f"{g.text(r, col)} | {g.text(r, col + 1)}",
                                        "Text in a row without a time label; not imported", date=d)
            prev_end = None
            for r, (start, stop) in time_rows:
                if prev_end is not None and start != prev_end:
                    self.issues.add("check", "time_label", cell_ref(g.title, r, 1), g.text(r, 1),
                                    f"Time label does not follow the previous row ({fmt_12(prev_end)})", date=d)
                prev_end = stop
                self.read_cell(d, g, r, col, start, stop)
            for r in staff_rows:
                self.read_staff(d, g, r, col, hdr)

    def read_cell(self, d, g, r, col, start, stop):
        sess = g.get(r, col)
        coach = g.get(r, col + 1)
        if isinstance(sess, (dt.date, dt.datetime)) or isinstance(coach, (dt.date, dt.datetime)):
            return
        if g.same_merge((r, col), (r, col + 1)):
            coach = None                           # one merged cell across both columns
        ref = cell_ref(g.title, r, col, col + 1)
        s_txt = "" if sess is None else str(sess)
        c_txt = "" if coach is None else str(coach)
        if not s_txt.strip():
            if c_txt.strip():
                it = self.classify(c_txt)
                if it.type == "holiday":
                    self.add_slot(d, start, stop, it, [], [], ref, s_txt, c_txt, [])
                else:
                    who, bad, _ = self.names.parse_coach_text(c_txt)
                    self.issues.add("check", "coach_without_session", ref, c_txt,
                                    "Coach column filled but the session column is empty; not imported", date=d)
            return
        pairs, how = pair_cell(s_txt, c_txt, self.classify, self.names.coach)
        flags = []
        if how == "merged":
            self.issues.add("check", "activity_merged", ref, f"{s_txt} || {c_txt}",
                            "Lines did not pair up; an activity line (e.g. 'Shooting') was joined to the "
                            "group line above it", suggestion="Check the coaches on these events", date=d)
        if how == "name_order":
            self.issues.add("check", "name_order", ref, f"{s_txt} || {c_txt}",
                            "Coach names were on one line; paired with the sessions in the order written",
                            suggestion="Put each coach on the same line as their session", date=d)
        if how == "ambiguous":
            flags.append("coach_pairing_unclear")
            self.issues.add("fix", "pairing", ref, f"{s_txt} || {c_txt}",
                            "Cannot tell which coach takes which session (line counts differ); "
                            "every coach was attached to every session in this cell",
                            suggestion="Put one session per line and the matching coach on the same line", date=d)
        for item_txt, coach_line in pairs:
            it = self.classify(item_txt)
            coaches, unresolved = [], []
            if coach_line:
                coaches, unresolved, junk = self.names.parse_coach_text(coach_line)
                if junk:
                    self.issues.add("check", "stray_characters", ref, coach_line,
                                    "Stray symbols in the coach cell were ignored", date=d)
                for u in unresolved:
                    sug = self.names.coach_suggestion(u)
                    as_session = self.classify(u)
                    if as_session.groups or as_session.activities:
                        prob = f"'{u}' in the coach column looks like a session, not a person"
                    else:
                        prob = f"Unknown coach name '{u}'"
                    self.issues.add("fix", "unknown_coach", ref, coach_line, prob,
                                    suggestion=("Did you mean " + " or ".join(sug) + "? Add to aliases.csv if so")
                                    if sug else "Add to aliases.csv as a coach, or fix the cell", date=d)
            if it.leftover:
                self.issues.add("check", "unparsed_text", ref, item_txt,
                                "Could not read: " + ", ".join(f"'{x}'" for x in it.leftover), date=d)
            if it.type == "other" and not it.activities:
                self.issues.add("fix", "unknown_session", ref, item_txt,
                                "Session not recognised; imported as type 'other'", date=d)
            if it.type == "pt" and not it.groups and not it.athletes:
                self.issues.add("check", "pt_without_group", ref, item_txt,
                                "PT with no group number or player named", date=d)
            athletes = [self.names.athlete(a, "roster grid") for a in it.athletes]
            self.add_slot(d, start, stop, it, coaches, athletes, ref, s_txt, c_txt, flags, unresolved)

    def add_slot(self, d, start, stop, it, coaches, athletes, ref, s_txt, c_txt, flags, unresolved=()):
        self.slots.append({
            "date": d, "start": start, "end": stop, "type": it.type, "title": it.title,
            "groups": list(it.groups), "coaches": list(coaches), "coaches_unresolved": list(unresolved),
            "athletes": list(athletes), "notes": list(it.notes), "cells": [ref],
            "raw_session": s_txt, "raw_coach": c_txt, "flags": list(flags),
        })

    def read_staff(self, d, g, r, col, hdr):
        raw_name = g.text(r, 1)
        who, bad, _ = self.names.parse_coach_text(raw_name)
        v = g.get(r, col + 1)
        ref = cell_ref(g.title, r, col + 1)
        if not who:
            if v not in (None, "") and not isinstance(v, (dt.date, dt.datetime)):
                self.issues.add("fix", "unknown_coach", cell_ref(g.title, r, 1), raw_name,
                                f"Hours row for unknown person '{norm_space(raw_name)}'",
                                suggestion="Add to aliases.csv as a coach", date=d)
            return
        coach = who[0]
        rec = {"person": coach, "date": d, "hours": None, "status": None, "note": "", "cell": ref}
        if v is None or (isinstance(v, str) and not v.strip()):
            return
        if isinstance(v, (int, float)):
            rec["hours"] = float(v)
            rec["status"] = "working"
        else:
            s = norm_space(v)
            key = s.upper().replace(" ", "")
            m = re.fullmatch(r"(\d+(?:\.\d+)?)\s*(.*)", s)
            if key in STATUS:
                rec["status"] = STATUS[key]
            elif m:
                rec["hours"] = float(m.group(1))
                rec["status"] = "working"
                rec["note"] = m.group(2)
            else:
                rec["status"] = "unknown"
                rec["note"] = s
                self.issues.add("check", "staff_status", ref, s, f"Unrecognised hours/status for {coach}", date=d)
        self.staff[(coach, d)] = rec

    def check_weekly_totals(self):
        weeks = collections.defaultdict(list)
        for (coach, d), rec in self.staff.items():
            g, hdr, col = self.day_source[d]
            weeks[(coach, g.title, hdr)].append(rec)
        for (coach, tab, hdr), recs in weeks.items():
            g = self.grids[tab]
            row = None
            for r in range(hdr + 1, hdr + 60):
                who, _, _ = self.names.parse_coach_text(g.text(r, 1)) if isinstance(g.get(r, 1), str) and not parse_time_label(g.get(r, 1)) else ([], 0, 0)
                if who and who[0] == coach:
                    row = r
                    break
            if row is None:
                continue
            block = next((b for b in self.blocks if b[0] is g and b[1] == hdr), None)
            if block is None or len(block[3]) < 7:
                continue                          # part-week block: the total may cover other tabs
            total_col = block[4]
            total = g.get(row, total_col)
            if not isinstance(total, (int, float)):
                continue
            # sum across the whole week row, not only the days inside the date range
            s = 0.0
            for c in DAY_COLS:
                if c + 1 >= total_col:
                    break
                v = g.get(row, c + 1)
                if isinstance(v, (int, float)):
                    s += float(v)
                elif isinstance(v, str):
                    m = re.match(r"\s*(\d+(?:\.\d+)?)", v)
                    if m:
                        s += float(m.group(1))
            if abs(s - float(total)) > 0.01:
                self.issues.add("check", "hours_total", cell_ref(tab, row, total_col), str(total),
                                f"{coach}'s weekly total says {total:g} h but the days add up to {s:g} h")

    # -- side tables ------------------------------------------------------
    def read_side_tables(self):
        group_rx = re.compile(r"^\s*(pt\s*(sr|jr)\s*\d|pt\s*rehab|pt|asb\s*(seniors?|juniors?|jr|sr|rehab|u18))\s*$", re.I)
        for g, hdr, end, dates, _ in self.blocks:
            week = sorted(dates.values())
            if not week or not any(self.date_from <= d <= self.date_to for d in week):
                continue
            if not any(self.day_source.get(d, (None,))[0] is g for d in week):
                continue
            for r in range(hdr, end + 1):
                labels = {c: g.text(r, c) for c in range(SIDE_FIRST_COL, g.max_col + 1)
                          if group_rx.match(g.text(r, c) or "")}
                if len(labels) < 2:
                    continue
                self.read_one_table(g, r, labels, week)

    def read_one_table(self, g, r, labels, week):
        lo, hi = min(labels), max(labels)
        cols = range(max(SIDE_FIRST_COL, lo - 1), hi + 1)
        for c in cols:
            players = []
            rr = r + 1
            blanks = 0
            while rr <= g.max_row and rr <= r + 25:
                v = g.text(rr, c).strip()
                if not v:
                    blanks += 1
                    if blanks >= 2:
                        break
                    rr += 1
                    continue
                blanks = 0
                if not self.name_like(v):
                    break
                players.append((rr, v))
                rr += 1
            label = labels.get(c, "").strip()
            coach_raw = g.text(r - 1, c).strip() if r > 1 else ""
            if coach_raw.lower() == "coach name":
                coach_raw = ""
            if not players and not label:
                continue
            ref = cell_ref(g.title, r, c)
            if not label:
                self.issues.add("fix", "side_table", ref, ", ".join(p for _, p in players),
                                "Side-table column with players but no group name"
                                + (f" (coach above: {coach_raw})" if coach_raw else ""),
                                suggestion="Add the group name (e.g. PT Sr 1) above the players", date=week[0])
                continue
            group = self.side_group_name(label)
            if group == "PT":
                self.issues.add("fix", "side_table", ref, label,
                                "Group is labelled just 'PT' - which PT group is it?", date=week[0])
            if coach_raw:
                coaches, bad, _ = self.names.parse_coach_text(coach_raw)
                for cname in coaches:
                    self.memberships.append({"person": cname, "role": "coach", "group": group,
                                             "week_start": week[0], "week_end": week[-1],
                                             "cell": cell_ref(g.title, r - 1, c), "note": ""})
                for u in bad:
                    self.issues.add("fix", "unknown_coach", cell_ref(g.title, r - 1, c), coach_raw,
                                    f"Unknown coach name '{u}' above {group}", date=week[0])
            for rr, p in players:
                note = ""
                m = re.match(r"^(.*?)\s*\((.*)\)\s*$", p)
                if m:
                    p, note = m.group(1), m.group(2)
                name = self.names.athlete(p, "side table")
                self.memberships.append({"person": name, "role": "athlete", "group": group,
                                         "week_start": week[0], "week_end": week[-1],
                                         "cell": cell_ref(g.title, rr, c), "note": note})

    @staticmethod
    def name_like(v):
        if "\n" in v or len(v) > 40:
            return False
        if parse_time_label(v) or as_date(v):
            return False
        if not re.fullmatch(r"[A-Za-z][A-Za-z .'-]*(\([^)]*\))?", v.strip()):
            return False
        words = v.lower().split()
        bad = {"game", "rest", "court", "group", "cls", "pt", "snc", "session", "coach", "head",
               "team", "availability", "timings", "break", "morning", "for", "weeks", "days", "on",
               "m", "mon", "tu", "tue", "tues", "wed", "th", "thu", "thur", "thru", "thurs", "fri", "sat", "sun"}
        return len(words) <= 4 and not (set(words) & bad)

    @staticmethod
    def side_group_name(label):
        l = norm_space(label).lower()
        m = re.fullmatch(r"pt\s*(sr|jr)\s*(\d)", l)
        if m:
            return f"PT {m.group(1).title()} {m.group(2)}"
        if re.fullmatch(r"pt\s*rehab", l):
            return "PT Rehab"
        if l == "pt":
            return "PT"
        if re.search(r"rehab", l):
            return "ASB Rehab"
        if re.search(r"u18", l):
            return "ASB U18"
        if re.search(r"sen|sr", l):
            return "ASB Seniors"
        if re.search(r"jun|jr", l):
            return "ASB Juniors"
        return nice_case(label)

    # -- merge ------------------------------------------------------------
    def merge_slots(self):
        """Consecutive half-hours with the same session and coaches become one event."""
        by_day = collections.defaultdict(list)
        for s in self.slots:
            by_day[s["date"]].append(s)
        events = []
        for d in sorted(by_day):
            open_ = {}
            for s in sorted(by_day[d], key=lambda x: x["start"]):
                key = (s["type"], s["title"], tuple(sorted(s["coaches"])), tuple(sorted(s["coaches_unresolved"])))
                cur = open_.get(key)
                if cur and cur["end"] == s["start"]:
                    cur["end"] = s["end"]
                    for ref in s["cells"]:
                        if ref not in cur["cells"]:
                            cur["cells"].append(ref)
                    for f in s["flags"]:
                        if f not in cur["flags"]:
                            cur["flags"].append(f)
                else:
                    ev = dict(s, cells=list(s["cells"]), flags=list(s["flags"]))
                    open_[key] = ev
                    events.append(ev)
        seen = collections.Counter()
        for ev in events:
            base = f"{ev['date']:%Y%m%d}-{ev['start']:04d}-{slug(ev['title'])[:40]}"
            seen[base] += 1
            ev["id"] = base if seen[base] == 1 else f"{base}-{seen[base]}"
            ev["source_ref"] = merge_refs(ev["cells"])
        return events

    # -- cross checks -----------------------------------------------------
    def check_staff_conflicts(self, events):
        busy = collections.defaultdict(list)
        for ev in events:
            for c in ev["coaches"]:
                busy[(c, ev["date"])].append(ev)
        for (coach, d), rec in self.staff.items():
            if rec["status"] in ("WO", "L", "SL") and busy.get((coach, d)):
                evs = busy[(coach, d)]
                self.issues.add("check", "off_but_scheduled", rec["cell"], rec["status"],
                                f"{coach} is marked {rec['status']} on {d:%a %d %b} but is named on "
                                + ", ".join(f"{fmt_12(e['start'])} {e['title']}" for e in evs[:3])
                                + ("..." if len(evs) > 3 else ""), date=d)


def merge_refs(refs):
    """["'Sep'!B4:C4", "'Sep'!B5:C5"] -> "'Sep'!B4:C5"."""
    out = []
    for ref in refs:
        m = re.fullmatch(r"(.*)!([A-Z]+)(\d+)(?::([A-Z]+)(\d+))?", ref)
        if not m:
            out.append(ref)
            continue
        tab, c1, r1, c2, r2 = m.group(1), m.group(2), int(m.group(3)), m.group(4) or m.group(2), int(m.group(5) or m.group(3))
        if out and isinstance(out[-1], list) and out[-1][0] == tab and out[-1][1] == c1 and out[-1][3] == c2 and out[-1][4] + 1 == r1:
            out[-1][4] = r2
        else:
            out.append([tab, c1, r1, c2, r2])
    return ", ".join(x if isinstance(x, str) else f"{x[0]}!{x[1]}{x[2]}:{x[3]}{x[4]}" for x in out)


def merge_memberships(rows):
    """Weekly snapshots -> ranges. A run continues while the person is in the group in
    consecutive snapshots of that group; a snapshot without them ends the run."""
    snaps = collections.defaultdict(set)
    weeks_by_group = collections.defaultdict(set)
    notes = {}
    for m in rows:
        wk = (m["week_start"], m["week_end"])
        snaps[(m["person"], m["role"], m["group"])].add(wk)
        weeks_by_group[m["group"]].add(wk)
        if m["note"]:
            notes[(m["person"], m["group"])] = m["note"]
    out = []
    for (person, role, group), weeks in snaps.items():
        order = sorted(weeks_by_group[group])
        run = None
        for wk in order:
            if wk in weeks:
                if run is None:
                    run = [wk[0], wk[1], 1]
                else:
                    run[1] = wk[1]
                    run[2] += 1
            elif run:
                out.append((person, role, group, run))
                run = None
        if run:
            out.append((person, role, group, run))
    res = [{"person": p, "role": r, "group": g, "from": run[0], "to": run[1], "snapshots": run[2],
            "note": notes.get((p, g), "")} for p, r, g, run in out]
    res.sort(key=lambda x: (x["group"], x["role"] != "coach", x["person"], x["from"]))
    return res


# ---------------------------------------------------------------- output
def write_csv(path, rows, fields):
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
        w.writeheader()
        for r in rows:
            w.writerow({k: (", ".join(v) if isinstance(v, list) else v) for k, v in r.items()})


def build_model(imp, events, memberships):
    people = {}
    for c in COACHES:
        people[f"coach:{slug(c)}"] = {"id": f"coach:{slug(c)}", "display_name": c, "role": "coach",
                                     "aliases": sorted({s for s in COACHES[c]})}
    for a, rec in sorted(imp.names.athletes.items()):
        pid = f"athlete:{slug(a)}"
        people[pid] = {"id": pid, "display_name": a, "role": "athlete",
                       "aliases": sorted(s for s in rec["spellings"] if s != a)}
    groups = {}
    for ev in events:
        for gname in ev["groups"]:
            groups.setdefault(gname, {"id": f"group:{slug(gname)}", "name": gname, "kind": group_kind(gname)})
    for m in memberships:
        groups.setdefault(m["group"], {"id": f"group:{slug(m['group'])}", "name": m["group"], "kind": group_kind(m["group"])})
    groups.setdefault(ALL_COACHES, {"id": "group:all-coaches", "name": ALL_COACHES, "kind": "staff"})

    def pid(name, role):
        return f"{role}:{slug(name)}"

    ev_rows, eg_rows, ep_rows = [], [], []
    for ev in events:
        ev_rows.append({"id": ev["id"], "date": str(ev["date"]), "start": fmt_time(ev["start"]),
                        "end": fmt_time(ev["end"]), "type": ev["type"], "title": ev["title"],
                        "location": "", "source_ref": ev["source_ref"], "needs_review": bool(ev["flags"]),
                        "flags": ev["flags"]})
        for gname in ev["groups"]:
            eg_rows.append({"event": ev["id"], "group": groups[gname]["id"]})
        for c in ev["coaches"]:
            if c == ALL_COACHES:
                eg_rows.append({"event": ev["id"], "group": "group:all-coaches"})
            else:
                ep_rows.append({"event": ev["id"], "person": pid(c, "coach"), "role": "coach"})
        for u in ev["coaches_unresolved"]:
            ep_rows.append({"event": ev["id"], "person": None, "raw_name": u, "role": "coach_unresolved"})
        for a in ev["athletes"]:
            ep_rows.append({"event": ev["id"], "person": pid(a, "athlete"), "role": "athlete"})
    ms = [{"person": pid(m["person"], m["role"]), "group": groups[m["group"]]["id"], "role": m["role"],
           "from": str(m["from"]), "to": str(m["to"]), "note": m["note"]} for m in memberships]
    st = [{"person": pid(r["person"], "coach"), "date": str(r["date"]), "status": r["status"],
           "hours": r["hours"], "note": r["note"], "source_ref": r["cell"]}
          for r in sorted(imp.staff.values(), key=lambda x: (x["date"], x["person"]))]
    for c in {r["person"] for r in imp.staff.values()}:
        people.setdefault(pid(c, "coach"), {"id": pid(c, "coach"), "display_name": c, "role": "coach", "aliases": []})
    return {"person": list(people.values()), "group": list(groups.values()), "membership": ms,
            "event": ev_rows, "event_group": eg_rows, "event_person": ep_rows, "staff_status": st,
            "import_issue": imp.issues.rows}


TYPE_LABEL = {"session": "Session", "pt": "PT", "match": "Match", "review": "Film", "ice_bath": "Ice bath",
              "physio": "Physio", "sc": "S&C", "meeting": "Meeting", "tournament": "Tournament",
              "holiday": "Holiday", "other": "Other"}
KIND_LABEL = {
    "pairing": "Coach and session lines don't pair up",
    "unknown_coach": "Unknown coach names",
    "unknown_session": "Sessions not recognised",
    "unparsed_text": "Text the importer could not read",
    "header_date": "Week header dates",
    "duplicate_day": "Same day in more than one tab",
    "side_table": "PT / squad side-tables",
    "coach_without_session": "Coach named but no session",
    "no_time_label": "Rows without a time label",
    "time_label": "Time labels",
    "pt_without_group": "PT with no group or player",
    "stray_characters": "Stray symbols in coach cells",
    "staff_status": "Hours / status cells",
    "hours_total": "Weekly hour totals that don't add up",
    "off_but_scheduled": "Coach marked off but on the roster",
    "activity_merged": "Activity line joined to its group line",
    "name_order": "Coaches on one line, paired in written order",
}


def render_report(imp, events, memberships, model, args, generated):
    e = html.escape
    issues = imp.issues.rows
    sev_count = collections.Counter(i["severity"] for i in issues)
    by_kind = collections.defaultdict(list)
    for i in issues:
        by_kind[i["kind"]].append(i)
    days = sorted({ev["date"] for ev in events})
    type_count = collections.Counter(ev["type"] for ev in events)
    coach_names = sorted({c for ev in events for c in ev["coaches"] if c != ALL_COACHES})
    athletes = imp.names.athletes
    clusters = imp.names.similar_athletes()

    H = []
    H.append("""<!doctype html><html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>Roster import review</title>
<style>
:root{--bg:#f4f6f4;--card:#fff;--ink:#161c21;--muted:#5f6b66;--rule:#dbe1dc;--accent:#0e6a58;
--fix:#a2331b;--fixbg:#f8e6e1;--check:#8a5a00;--checkbg:#f7efdc;--soft:#eef2ee}
@media (prefers-color-scheme:dark){:root{--bg:#111514;--card:#1a201e;--ink:#e6ebe8;--muted:#9aa7a1;
--rule:#2c3532;--accent:#5cc3a8;--fix:#ff9c85;--fixbg:#3a211b;--check:#f0c060;--checkbg:#342a14;--soft:#222a27}}
*{box-sizing:border-box}body{margin:0;background:var(--bg);color:var(--ink);
font:15px/1.5 -apple-system,BlinkMacSystemFont,"Segoe UI",Roboto,Arial,sans-serif}
main{max-width:980px;margin:0 auto;padding:16px}
h1{font-size:22px;margin:8px 0 4px}h2{font-size:18px;margin:28px 0 8px}h3{font-size:15px;margin:16px 0 6px}
p{margin:6px 0}.muted{color:var(--muted)}.card{background:var(--card);border:1px solid var(--rule);
border-radius:10px;padding:12px 14px;margin:10px 0}
.stats{display:grid;grid-template-columns:repeat(auto-fit,minmax(130px,1fr));gap:8px}
.stat{background:var(--card);border:1px solid var(--rule);border-radius:10px;padding:10px 12px}
.stat b{display:block;font-size:22px}.stat span{color:var(--muted);font-size:13px}
table{border-collapse:collapse;width:100%;font-size:13px}th,td{text-align:left;padding:6px 8px;
border-bottom:1px solid var(--rule);vertical-align:top}th{color:var(--muted);font-weight:600}
.scroll{overflow-x:auto}.pill{display:inline-block;padding:1px 8px;border-radius:99px;font-size:12px;
background:var(--soft);white-space:nowrap}.fix{background:var(--fixbg);color:var(--fix)}
.check{background:var(--checkbg);color:var(--check)}code,.mono{font-family:ui-monospace,Menlo,Consolas,monospace;font-size:12px}
details{background:var(--card);border:1px solid var(--rule);border-radius:10px;margin:8px 0}
summary{cursor:pointer;padding:10px 14px;font-weight:600}details>div{padding:0 14px 12px}
.ev{display:grid;grid-template-columns:92px 1fr;gap:8px;padding:6px 0;border-bottom:1px solid var(--rule)}
.ev:last-child{border-bottom:0}.flag{color:var(--fix);font-weight:600}a{color:var(--accent)}
ol li,ul li{margin:4px 0}
</style></head><body><main>""")
    H.append(f"<h1>Roster 2026 &rarr; app: import review</h1>")
    H.append(f"<p class=muted>{args.date_from:%d %b} &ndash; {args.date_to:%d %b %Y} &middot; generated {e(generated)} "
             f"&middot; source <a href='{SHEET_URL}'>Roster 2026</a> (read-only)</p>")
    H.append("<div class=stats>")
    for n, label in [(len(events), "events"), (len(days), "days with events"), (len(coach_names), "coaches named"),
                     (len(athletes), "athlete names"), (sev_count.get("fix", 0), "need fixing"),
                     (sev_count.get("check", 0), "to check")]:
        H.append(f"<div class=stat><b>{n}</b><span>{label}</span></div>")
    H.append("</div>")

    H.append("<div class=card><b>How to read this.</b> The importer only reads the sheet. Every half-hour "
             "cell was split into sessions, paired with coaches, and merged into events. Anything it could "
             "not resolve is listed below with the exact sheet cell, instead of being guessed. "
             "<span class='pill fix'>fix</span> = the calendar will be wrong until this is fixed in the sheet "
             "or confirmed. <span class='pill check'>check</span> = probably fine, worth a look.</div>")

    # decisions
    H.append("<h2>Decisions needed before the calendar is built</h2><div class=card><ol>")
    H.append("<li><b>Sign off this import</b> (or list what is wrong) &mdash; the next steps build on it.</li>")
    grp_counts = collections.Counter(g for ev in events for g in ev["groups"])
    if grp_counts.get("Elite") and grp_counts.get("ASB Seniors"):
        H.append(f"<li><b>Is &ldquo;Elite&rdquo; the same squad as &ldquo;ASB Seniors&rdquo;</b>, and &ldquo;Jr Elite&rdquo; the same "
                 f"as &ldquo;ASB Juniors&rdquo;? The sheet uses both ({grp_counts['Elite']} / {grp_counts['ASB Seniors']} events). "
                 "Kept separate until you say.</li>")
    if clusters:
        H.append(f"<li><b>{len(clusters)} sets of athlete names may be the same person</b> (e.g. "
                 + e(", ".join(" / ".join(c) for c in clusters[:3])) + "). Confirm in the table below; "
                 "confirmed ones go into <code>aliases.csv</code>.</li>")
    unk = sorted({i['raw'] for i in by_kind.get('unknown_coach', [])})
    if unk:
        H.append(f"<li><b>{len(by_kind['unknown_coach'])} cells name a coach the importer does not know</b> &mdash; "
                 "typos or new staff. See the list below.</li>")
    H.append("<li><b>Group sessions have no athlete list.</b> The sheet names athletes only for PTs and "
             "individuals, so ASB Seniors / Intermediate / Mini sessions show coaches only until the "
             "membership table (step 2) is filled in.</li>")
    H.append("</ol></div>")

    # issues
    H.append("<h2>Issues</h2>")
    order = ["pairing", "unknown_coach", "unknown_session", "side_table", "header_date", "duplicate_day",
             "activity_merged", "name_order", "unparsed_text", "pt_without_group", "coach_without_session", "no_time_label", "time_label",
             "off_but_scheduled", "hours_total", "staff_status", "stray_characters"]
    for kind in order + [k for k in by_kind if k not in order]:
        rows = by_kind.get(kind)
        if not rows:
            continue
        sev = "fix" if any(r["severity"] == "fix" for r in rows) else "check"
        # the same text often repeats in consecutive half-hours: show it once with its cells
        grouped = collections.OrderedDict()
        for r in sorted(rows, key=lambda x: (x["date"], x["cell"])):
            grouped.setdefault((r["raw"], r["problem"], r["suggestion"]), []).append(r)
        H.append(f"<details{' open' if sev == 'fix' and len(grouped) <= 12 else ''}><summary>"
                 f"<span class='pill {sev}'>{sev}</span> {e(KIND_LABEL.get(kind, kind))} &middot; "
                 f"{len(rows)} cell{'s' if len(rows) != 1 else ''}</summary><div class=scroll><table>"
                 "<tr><th>Sheet text</th><th>Problem</th><th>Where</th><th>Suggestion</th></tr>")
        for (raw, problem, suggestion), rs in grouped.items():
            where = "<br>".join(f"<span class=mono>{e(x['date'])} {e(x['cell'])}</span>" for x in rs[:4])
            if len(rs) > 4:
                where += f"<br><span class=muted>+{len(rs) - 4} more (issues.csv)</span>"
            H.append(f"<tr><td>{e(raw) or '<span class=muted>(empty)</span>'}</td><td>{e(problem)}</td>"
                     f"<td>{where}</td><td>{e(suggestion)}</td></tr>")
        H.append("</table></div></details>")

    # names
    H.append("<h2>Names</h2>")
    H.append("<h3>Coaches: every spelling seen</h3><div class='card scroll'><table><tr><th>Coach</th><th>Spellings in the sheet</th><th>Cells</th></tr>")
    per_coach = collections.defaultdict(collections.Counter)
    for (raw, res, role), n in imp.names.seen.items():
        if role == "coach":
            per_coach[res or "(unresolved)"][raw] += n
    for coach in sorted(per_coach, key=lambda x: (x == "(unresolved)", x)):
        sp = per_coach[coach]
        H.append(f"<tr><td><b>{e(coach)}</b></td><td>{e(', '.join(f'{k} ({v})' for k, v in sp.most_common()))}</td>"
                 f"<td>{sum(sp.values())}</td></tr>")
    H.append("</table></div>")
    if clusters:
        H.append("<h3>Athlete names that might be one person</h3><div class='card scroll'><table><tr><th>Names</th><th>Seen</th></tr>")
        for cl in clusters:
            H.append("<tr><td>" + e(" / ".join(cl)) + "</td><td>" +
                     e(", ".join(f"{n}: {sum(athletes[n]['sources'].values())}" for n in cl)) + "</td></tr>")
        H.append("</table><p class=muted>Not merged. Add a line per confirmed spelling to aliases.csv, e.g. "
                 "<code>Asha K,Asha Kumar,athlete</code>.</p></div>")
    H.append("<h3>All athlete names</h3><div class='card scroll'><table><tr><th>Name</th><th>Where seen</th><th>Times</th></tr>")
    for a in sorted(athletes):
        rec = athletes[a]
        H.append(f"<tr><td>{e(a)}</td><td>{e(', '.join(rec['sources']))}</td><td>{sum(rec['sources'].values())}</td></tr>")
    H.append("</table></div>")

    # groups
    H.append("<h2>Groups found</h2><div class='card scroll'><table><tr><th>Group</th><th>Kind</th><th>Events</th></tr>")
    for gname, n in grp_counts.most_common():
        H.append(f"<tr><td>{e(gname)}</td><td>{e(group_kind(gname))}</td><td>{n}</td></tr>")
    H.append("</table></div>")

    # memberships
    H.append("<h2>Membership seed (from the PT / squad side-tables)</h2>")
    H.append("<p class=muted>Each side-table is a weekly snapshot. A person stays in a group from the first to the last "
             "consecutive snapshot they appear in. Weeks without a side-table are not counted either way.</p>")
    by_group = collections.defaultdict(list)
    for m in memberships:
        by_group[m["group"]].append(m)
    H.append("<div class='card scroll'><table><tr><th>Group</th><th>Coach(es)</th><th>Athletes</th></tr>")
    for gname in sorted(by_group):
        ms = by_group[gname]
        co = [f"{m['person']} ({m['from']:%d %b}&ndash;{m['to']:%d %b})" for m in ms if m["role"] == "coach"]
        at = [f"{m['person']}" + (f" [{m['note']}]" if m['note'] else "") + f" ({m['from']:%d %b}&ndash;{m['to']:%d %b})"
              for m in ms if m["role"] == "athlete"]
        H.append(f"<tr><td><b>{e(gname)}</b></td><td>{'<br>'.join(e(x).replace('&amp;ndash;', '&ndash;') for x in co)}</td>"
                 f"<td>{'<br>'.join(e(x).replace('&amp;ndash;', '&ndash;') for x in at)}</td></tr>")
    H.append("</table></div>")

    # staff
    H.append("<h2>Coach hours and days off</h2><div class='card scroll'><table><tr><th>Coach</th><th>Days worked</th>"
             "<th>Hours</th><th>WO</th><th>Leave (L)</th><th>Sick (SL)</th></tr>")
    agg = collections.defaultdict(lambda: collections.Counter())
    for rec in imp.staff.values():
        a = agg[rec["person"]]
        if rec["status"] == "working":
            a["days"] += 1
            a["hours"] += rec["hours"] or 0
        elif rec["status"]:
            a[rec["status"]] += 1
    for coach in sorted(agg):
        a = agg[coach]
        H.append(f"<tr><td>{e(coach)}</td><td>{a['days']}</td><td>{a['hours']:g}</td><td>{a['WO']}</td><td>{a['L']}</td><td>{a['SL']}</td></tr>")
    H.append("</table></div>")

    # skipped tabs
    H.append("<h2>Tabs not imported</h2><div class=card><ul>")
    for t, why in imp.skipped_tabs:
        H.append(f"<li><b>{e(t.strip())}</b> &mdash; {e(why)}</li>")
    H.append("</ul></div>")

    # day by day
    H.append("<h2>Day by day</h2><p class=muted>Every imported event. Red = needs review.</p>")
    evs_by_day = collections.defaultdict(list)
    for ev in events:
        evs_by_day[ev["date"]].append(ev)
    for d in sorted(set(list(evs_by_day) + list(imp.day_source))):
        evs = sorted(evs_by_day.get(d, []), key=lambda x: (x["start"], x["end"]))
        g, hdr, col = imp.day_source.get(d, (None, 0, 0))
        n_iss = sum(1 for i in issues if i["date"] == str(d))
        off = [f"{r['person']} {r['status']}" for (c, dd), r in imp.staff.items() if dd == d and r["status"] in ("WO", "L", "SL")]
        H.append(f"<details><summary>{d:%a %d %b} &middot; {len(evs)} events"
                 + (f" &middot; <span class='pill fix'>{n_iss} issues</span>" if n_iss else "") + "</summary><div>")
        if g:
            H.append(f"<p class=muted>From '{e(g.title.strip())}' tab, column {col_letter(col)}, week starting row {hdr}."
                     + (f" Off: {e(', '.join(sorted(off)))}." if off else "") + "</p>")
        for ev in evs:
            who = ", ".join(ev["coaches"] + [f"?{u}" for u in ev["coaches_unresolved"]]) or "no coach named"
            flag = " <span class=flag>check coaches</span>" if ev["flags"] else ""
            H.append(f"<div class=ev><div class=mono>{fmt_12(ev['start'])}&ndash;{fmt_12(ev['end'])}</div><div>"
                     f"<span class=pill>{e(TYPE_LABEL.get(ev['type'], ev['type']))}</span> <b>{e(ev['title'])}</b>{flag}<br>"
                     f"<span class=muted>{e(who)} &middot; <span class=mono>{e(ev['source_ref'])}</span></span></div></div>")
        H.append("</div></details>")
    H.append("<p class=muted style='margin-top:24px'>Generated by asb/roster_importer/import_roster.py. "
             "The sheet was not modified.</p></main></body></html>")
    return "".join(H)


# ---------------------------------------------------------------- main
def main(argv=None):
    ap = argparse.ArgumentParser(description="ASB Roster 2026 importer (read-only)")
    ap.add_argument("source")
    ap.add_argument("--from", dest="date_from", default="2026-06-01", type=dt.date.fromisoformat)
    ap.add_argument("--to", dest="date_to", default="2026-09-30", type=dt.date.fromisoformat)
    ap.add_argument("--aliases", default=None)
    ap.add_argument("--out", default="roster_import_out")
    args = ap.parse_args(argv)

    wb = load_workbook(args.source)
    issues = Issues()
    names = Names(args.aliases)
    imp = Importer(wb, args.date_from, args.date_to, names, issues)
    imp.find_blocks()
    imp.choose_day_sources()
    imp.scan_names()
    imp.read_days()
    imp.check_weekly_totals()
    imp.read_side_tables()
    events = imp.merge_slots()
    imp.check_staff_conflicts(events)
    memberships = merge_memberships(imp.memberships)
    model = build_model(imp, events, memberships)

    os.makedirs(args.out, exist_ok=True)
    generated = dt.datetime.now().strftime("%d %b %Y %H:%M")
    with open(os.path.join(args.out, "report.html"), "w", encoding="utf-8") as f:
        f.write(render_report(imp, events, memberships, model, args, generated))
    with open(os.path.join(args.out, "roster_import.json"), "w", encoding="utf-8") as f:
        json.dump(model, f, indent=1, default=str)
    write_csv(os.path.join(args.out, "events.csv"),
              [dict(ev, date=str(ev["date"]), start=fmt_time(ev["start"]), end=fmt_time(ev["end"])) for ev in events],
              ["id", "date", "start", "end", "type", "title", "groups", "coaches", "coaches_unresolved",
               "athletes", "notes", "flags", "source_ref", "raw_session", "raw_coach"])
    write_csv(os.path.join(args.out, "issues.csv"), issues.rows,
              ["severity", "kind", "date", "cell", "raw", "problem", "suggestion"])
    write_csv(os.path.join(args.out, "names.csv"),
              [{"raw": r, "resolved_to": res, "role": role, "count": n}
               for (r, res, role), n in sorted(names.seen.items(), key=lambda x: (x[0][2], x[0][1], x[0][0]))],
              ["raw", "resolved_to", "role", "count"])
    write_csv(os.path.join(args.out, "memberships.csv"),
              [dict(m, **{"from": str(m["from"]), "to": str(m["to"])}) for m in memberships],
              ["group", "role", "person", "from", "to", "snapshots", "note"])
    write_csv(os.path.join(args.out, "staff_status.csv"),
              [dict(r, date=str(r["date"])) for r in sorted(imp.staff.values(), key=lambda x: (x["date"], x["person"]))],
              ["date", "person", "status", "hours", "note", "cell"])

    sev = collections.Counter(i["severity"] for i in issues.rows)
    print(f"Imported {len(events)} events over {len({ev['date'] for ev in events})} days "
          f"({args.date_from} to {args.date_to}).")
    print(f"Issues: {sev.get('fix', 0)} to fix, {sev.get('check', 0)} to check.")
    print(f"Report: {os.path.join(args.out, 'report.html')}")


if __name__ == "__main__":
    main()
