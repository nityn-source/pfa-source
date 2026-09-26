"""Tests for the roster importer. Run:  python3 -m unittest test_import_roster.py

Uses a small made-up workbook (invented player names), never the real sheet.
"""
import csv
import datetime as dt
import os
import tempfile
import unittest

import openpyxl

import import_roster as ir


def classify(text):
    return ir.parse_session_item(text, set())


class TimeLabels(unittest.TestCase):
    def test_messy_labels(self):
        self.assertEqual(ir.parse_time_label("8:00AM- 8:30AM"), (480, 510))
        self.assertEqual(ir.parse_time_label("12:30PM - 1.00PM"), (750, 780))
        self.assertEqual(ir.parse_time_label("9:30AM - 10:AM"), (570, 600))
        self.assertEqual(ir.parse_time_label("7:00AM - 7:30 AM "), (420, 450))
        self.assertIsNone(ir.parse_time_label("Manoj"))


class Sessions(unittest.TestCase):
    def test_pt_groups(self):
        it = classify("Pt Jr 1 & PT Jr 2")
        self.assertEqual((it.type, it.groups), ("pt", ["PT Jr 1", "PT Jr 2"]))
        self.assertEqual(classify("PT JR 3 & PT Jr 1&4").groups, ["PT Jr 3", "PT Jr 1", "PT Jr 4"])
        self.assertEqual(classify("JR 3&4 PT").groups, ["PT Jr 3", "PT Jr 4"])

    def test_types_and_written_order(self):
        self.assertEqual(classify("SR ASB match").type, "match")
        self.assertEqual(classify("Ice Bath - SC").type, "ice_bath")
        self.assertEqual(classify("Coaches meeting").type, "meeting")
        self.assertEqual(classify("SC & Jr Elite").groups, ["SC", "Jr Elite"])

    def test_player_names(self):
        it = classify("Asha PT")
        self.assertEqual((it.type, it.athletes), ("pt", ["asha"]))
        it = classify("ASB Seniors     Asha assessment")
        self.assertEqual((it.type, it.groups, it.athletes), ("other", ["ASB Seniors"], ["asha"]))


class Pairing(unittest.TestCase):
    def test_line_for_line(self):
        pairs, how = ir.pair_cell("PT Sr 2\nPT Sr 3", "Manoj\nPuneeth", classify)
        self.assertEqual((pairs, how), ([("PT Sr 2", "Manoj"), ("PT Sr 3", "Puneeth")], "ok"))

    def test_wide_spaces(self):
        pairs, how = ir.pair_cell("ASB Seniors       Asha assessment", "Puneeth       Rudra sir", classify)
        self.assertEqual(how, "ok")
        self.assertEqual(pairs[1], ("Asha assessment", "Rudra sir"))

    def test_ice_bath_needs_no_coach(self):
        pairs, how = ir.pair_cell(" PT Sr 2           Ice Bath", "Manoj\n", classify)
        self.assertEqual((pairs, how), ([("PT Sr 2", "Manoj"), ("Ice Bath", None)], "ok"))

    def test_all_coaches(self):
        pairs, _ = ir.pair_cell("SC (Stations)\nmini", "ALL COACHES", classify)
        self.assertTrue(all(c == ir.ALL_COACHES for _, c in pairs))

    def test_ambiguous_is_flagged_not_guessed(self):
        _, how = ir.pair_cell("PT\nBeginner\nIntermediate", "Suraj\nManoj", classify)
        self.assertEqual(how, "ambiguous")


class CoachNames(unittest.TestCase):
    def setUp(self):
        self.names = ir.Names()

    def test_spellings(self):
        who, bad, _ = self.names.parse_coach_text("Rura sir, Divya ma'am / Surajj")
        self.assertEqual((who, bad), (["Rudra", "Divya", "Suraj"], []))

    def test_joined_and_parenthesis(self):
        self.assertEqual(self.names.parse_coach_text("Suraj Mario")[0], ["Suraj", "Mario"])
        self.assertEqual(self.names.parse_coach_text("ManojSanket")[0], ["Manoj", "Sanket"])
        self.assertEqual(self.names.parse_coach_text("Puneeth(Nityan B.)")[0], ["Puneeth", "Nityn"])

    def test_typo_is_not_merged(self):
        who, bad, _ = self.names.parse_coach_text("Marioo")
        self.assertEqual((who, bad), ([], ["Marioo"]))
        self.assertEqual(self.names.coach_suggestion("Marioo"), ["Mario"])


class EndToEnd(unittest.TestCase):
    def test_small_workbook(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sep"
        monday = dt.datetime(2026, 9, 7)
        for i, c in enumerate(ir.DAY_COLS):
            # Friday's header is a copy-paste typo; the importer must correct it
            ws.cell(1, c, monday + dt.timedelta(days=i if c != 10 else i - 7))
        ws.cell(2, 1, "10:00AM - 10:30AM")
        ws.cell(3, 1, "10:30AM - 11:00AM")
        for r in (2, 3):
            ws.cell(r, 2, "ASB Seniors")
            ws.cell(r, 3, "Puneeth")
        ws.cell(2, 4, "PT Sr 2\nAsha")
        ws.cell(2, 5, "Manoj\nRudra sir")
        ws.cell(4, 1, "Manoj")
        ws.cell(4, 3, 3)
        ws.cell(4, 5, "WO")
        ws.cell(4, 16, 3)
        with tempfile.TemporaryDirectory() as tmp:
            src = os.path.join(tmp, "roster.xlsx")
            wb.save(src)
            out = os.path.join(tmp, "out")
            ir.main([src, "--from", "2026-09-07", "--to", "2026-09-13", "--out", out])
            with open(os.path.join(out, "events.csv"), encoding="utf-8") as f:
                events = list(csv.DictReader(f))
            with open(os.path.join(out, "issues.csv"), encoding="utf-8") as f:
                issues = list(csv.DictReader(f))
        seniors = [e for e in events if e["title"] == "ASB Seniors"]
        self.assertEqual(len(seniors), 1)                        # two half-hours merged
        self.assertEqual((seniors[0]["start"], seniors[0]["end"]), ("10:00", "11:00"))
        pt = [e for e in events if e["title"] == "PT: Asha"][0]
        self.assertEqual(pt["coaches"], "Rudra")
        kinds = {i["kind"] for i in issues}
        self.assertIn("header_date", kinds)                     # Friday typo reported
        self.assertIn("off_but_scheduled", kinds)               # Manoj WO but on PT Sr 2


if __name__ == "__main__":
    unittest.main()
